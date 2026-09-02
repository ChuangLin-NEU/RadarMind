"""Chinese project-progress PPT. Numbers only from metric pack."""
from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import nsmap
from pptx.util import Emu, Inches, Pt

from .load_metrics import D, pct

NAVY = RGBColor(0x1B, 0x36, 0x5D)
BRICK = RGBColor(0xB4, 0x53, 0x2A)
INK = RGBColor(0x1F, 0x29, 0x37)
MUTED = RGBColor(0x5B, 0x64, 0x75)
BG = RGBColor(0xF7, 0xF4, 0xEE)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LINE = RGBColor(0xC9, 0xC1, 0xB2)


def _set_fill(shape, rgb):
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb
    shape.line.fill.background()


def _box(slide, l, t, w, h, fill=None):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(l), Inches(t), Inches(w), Inches(h))
    if fill is None:
        sh.fill.background()
        sh.line.fill.background()
    else:
        _set_fill(sh, fill)
    return sh


def _text(slide, l, t, w, h, text, *, size=18, bold=False, color=INK, align=PP_ALIGN.LEFT, font="Microsoft YaHei"):
    box = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font
    return box


def _bullets(slide, l, t, w, h, items, *, size=18, color=INK):
    box = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.space_after = Pt(8)
        run = p.add_run()
        run.text = item
        run.font.size = Pt(size)
        run.font.color.rgb = color
        run.font.name = "Microsoft YaHei"
        run.font.bold = False
    return box


def _notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


def _bg(slide):
    _box(slide, 0, 0, 13.333, 7.5, BG)
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(0.12), Inches(7.5))
    _set_fill(bar, NAVY)


def _title(slide, text):
    _text(slide, 0.5, 0.28, 12.4, 0.9, text, size=26, bold=True, color=NAVY)


def _footer(slide, n, ntot=15):
    _text(slide, 11.4, 7.12, 1.6, 0.28, f"{n} / {ntot}", size=11, color=MUTED, align=PP_ALIGN.RIGHT)


def _pic(slide, path: Path, l, t, w):
    slide.shapes.add_picture(str(path), Inches(l), Inches(t), width=Inches(w))


def build_ppt(pack: dict, fig_dir: Path, out_path: Path) -> dict:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]
    d = lambda k: D(pack, k)
    p = lambda k: pct(pack, k)

    # 1
    s = prs.slides.add_slide(blank)
    _box(s, 0, 0, 13.333, 7.5, NAVY)
    _text(s, 0.8, 2.1, 11.5, 1.2, "RadarMind", size=48, bold=True, color=WHITE)
    _text(s, 0.8, 3.4, 11.5, 1.2, "从“识别动作”到“理解不同雷达看到的同一种运动”", size=22, color=RGBColor(0xE8, 0xDF, 0xC8))
    _text(s, 0.8, 6.4, 10, 0.4, "跨雷达人体运动智能理解　项目进展", size=14, color=RGBColor(0xB7, 0xC4, 0xD6))
    _notes(
        s,
        "今天不讲软件功能，只讲一件事：我们一开始把雷达人体动作识别当成分类问题，后来发现真正卡住的是测量本身。"
        "听众如果没做过雷达也没关系，后面会先用一页把雷达怎么看见动作讲清楚。",
    )

    # 2
    s = prs.slides.add_slide(blank)
    _bg(s)
    _title(s, "雷达没有摄像头，它怎么知道人在走路还是弯腰？")
    _bullets(
        s,
        0.55,
        1.35,
        7.2,
        4.8,
        [
            "人一动，身体各部位的速度就不一样。",
            "雷达看到的不是照片，而是距离、速度、回波强弱随时间怎么变。",
            "这些变化叠在一起，形成这个动作的雷达指纹。",
        ],
        size=20,
    )
    # schematic
    boxes = [("人在动", 8.3, 2.0), ("雷达", 10.0, 2.0), ("时间–速度图", 8.3, 4.2), ("动作标签", 10.0, 4.2)]
    for name, x, y in boxes:
        sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(1.7), Inches(0.7))
        _set_fill(sh, WHITE)
        sh.line.color.rgb = NAVY
        tf = sh.text_frame
        tf.paragraphs[0].alignment = PP_ALIGN.CENTER
        r = tf.paragraphs[0].add_run()
        r.text = name
        r.font.size = Pt(14)
        r.font.color.rgb = NAVY
        r.font.name = "Microsoft YaHei"
        r.font.bold = True
    _text(s, 8.3, 6.3, 4.5, 0.4, "示意，不是实测图", size=12, color=MUTED)
    _footer(s, 2)
    _notes(
        s,
        "雷达没有像素人脸。人一动，胸、手、腿的径向速度不同，回波强弱也不同。"
        "把这些随时间排开，得到的是一张时间–速度图，不是一张照片。"
        "后面所有实验，都是在问：这张图里到底有多少是动作，有多少是这台雷达自己的痕迹。",
    )

    # 3
    s = prs.slides.add_slide(blank)
    _bg(s)
    _title(s, "只在同一台雷达上测试，模型其实已经很准了")
    nums = [("10 GHz", p("R0_10_WITHIN")), ("24 GHz", p("R0_24_WITHIN")), ("77 GHz", p("R0_77_WITHIN"))]
    for i, (lab, val) in enumerate(nums):
        x = 0.7 + i * 4.1
        card = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(2.0), Inches(3.7), Inches(3.0))
        _set_fill(card, WHITE)
        card.line.color.rgb = LINE
        _text(s, x, 2.25, 3.7, 0.4, lab, size=16, color=MUTED, align=PP_ALIGN.CENTER)
        _text(s, x, 2.8, 3.7, 1.2, val, size=44, bold=True, color=NAVY, align=PP_ALIGN.CENTER)
        _text(s, x, 4.2, 3.7, 0.4, "同雷达识别", size=14, color=MUTED, align=PP_ALIGN.CENTER)
    _text(s, 0.7, 5.5, 12, 0.8, "如果只看这些结果，很容易认为雷达动作识别已经做得不错。", size=18, color=INK)
    _footer(s, 3)
    _notes(
        s,
        f"10、24、77 GHz 各自训练、各自测试，准确率分别是 {p('R0_10_WITHIN')}、{p('R0_24_WITHIN')}、{p('R0_77_WITHIN')}。"
        "我们一开始也觉得分类器已经够用了。这一页请先记住对角线：同雷达看起来没问题。",
    )

    # 4
    s = prs.slides.add_slide(blank)
    _bg(s)
    _title(s, "换一台雷达，模型还认识这个动作吗？")
    left = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.6), Inches(1.7), Inches(5.6), Inches(3.6))
    _set_fill(left, WHITE)
    left.line.color.rgb = LINE
    _text(s, 0.6, 2.0, 5.6, 0.4, "同雷达", size=16, color=MUTED, align=PP_ALIGN.CENTER)
    _text(s, 0.6, 2.6, 5.6, 1.4, "接近 90%", size=48, bold=True, color=NAVY, align=PP_ALIGN.CENTER)
    _text(s, 0.6, 4.3, 5.6, 0.5, f"{p('R0_10_WITHIN')} / {p('R0_24_WITHIN')} / {p('R0_77_WITHIN')}", size=16, color=MUTED, align=PP_ALIGN.CENTER)
    right = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.1), Inches(1.7), Inches(5.6), Inches(3.6))
    _set_fill(right, RGBColor(0xF8, 0xEB, 0xE3))
    right.line.color.rgb = BRICK
    _text(s, 7.1, 2.0, 5.6, 0.4, "24 GHz → 77 GHz", size=16, color=BRICK, align=PP_ALIGN.CENTER)
    _text(s, 7.1, 2.6, 5.6, 1.4, p("R0_24_TO_77"), size=48, bold=True, color=BRICK, align=PP_ALIGN.CENTER)
    _text(s, 7.1, 4.3, 5.6, 0.5, "接近失效", size=16, color=BRICK, align=PP_ALIGN.CENTER)
    _text(s, 0.6, 5.7, 12, 0.7, "这成为整个项目的转折点：问题从“模型够不够大”变成了“它到底学到了什么”。", size=18, color=INK)
    if (fig_dir / "fig2_within_cross.png").is_file():
        pass  # numbers already on slide; matrix would crowd
    _footer(s, 4)
    _notes(
        s,
        f"把 24 GHz 上训好的模型直接放到 77 GHz，准确率只剩 {p('R0_24_TO_77')}。"
        "同雷达接近九成，换雷达接近失效。从这一页开始，我们不再先问怎么把网络做大，而是问它到底在认动作，还是在认这台雷达。",
    )

    # 5
    s = prs.slides.add_slide(blank)
    _bg(s)
    _title(s, "它学到的是“运动”，还是“这台雷达长什么样”？")
    _bullets(
        s,
        0.55,
        1.4,
        8.6,
        3.6,
        [
            "如果主要学的是人体运动：同一个 Bend，换雷达仍应像 Bend。",
            "实际：同一动作在不同雷达上的表示差得很远。",
            "口号可以先立在这里：Learn Motion, Not Sensor。",
        ],
        size=20,
    )
    stamp = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(9.3), Inches(2.2), Inches(3.5), Inches(2.4))
    _set_fill(stamp, NAVY)
    _text(s, 9.3, 2.45, 3.5, 0.5, "Learn Motion,", size=16, color=WHITE, align=PP_ALIGN.CENTER)
    _text(s, 9.3, 2.95, 3.5, 0.5, "Not Sensor", size=16, color=WHITE, align=PP_ALIGN.CENTER)
    _text(s, 9.3, 3.7, 3.5, 0.6, "Current status:\nNOT ACHIEVED", size=16, bold=True, color=RGBColor(0xF3, 0xD3, 0x9E), align=PP_ALIGN.CENTER)
    _footer(s, 5)
    _notes(
        s,
        "如果表示真的以人体运动为主，同一个弯腰换一台雷达，特征还应该落在弯腰附近。"
        "可三台雷达给出的图差得很大。Learn Motion, Not Sensor 是目标，不是已经做到的结果。右下角必须说清楚：现在还没做到。",
    )

    # 6
    s = prs.slides.add_slide(blank)
    _bg(s)
    _title(s, "第一嫌疑：三台雷达生成的“图”根本长得不一样")
    _bullets(
        s,
        0.55,
        1.3,
        12.2,
        3.2,
        [
            "官方图：尺寸不同，绘图方式不同，色条和边框也不同。",
            f"把明显绘图差异统一以后，10→24：{p('R0_10_TO_24')} → {p('R1_10_TO_24')}。",
            f"雷达身份识别仍然是 {p('RID_R0_3CLASS')}。",
            "图片处理确实有一部分责任，但不是全部原因。",
        ],
        size=20,
    )
    _footer(s, 6)
    _notes(
        s,
        f"三台雷达的官方图本来就不是同一种画法。裁掉色条和边框之后，10 到 24 GHz 从 {p('R0_10_TO_24')} 升到 {p('R1_10_TO_24')}，说明绘图方式确实掺了一脚。"
        f"可是用一张图判断“这是哪台雷达”，准确率仍是 {p('RID_R0_3CLASS')}。所以这一层只能算部分原因。",
    )

    # 7
    s = prs.slides.add_slide(blank)
    _bg(s)
    _title(s, "不看 PNG 了，直接从原始雷达数据重新算")
    _text(s, 0.55, 1.35, 12, 0.5, "RAW IQ  →  信号处理  →  数值 Doppler  →  S(时间, Doppler)", size=20, bold=True, color=NAVY)
    _bullets(
        s,
        0.55,
        2.1,
        12,
        3.4,
        [
            "这一步拿掉坐标轴、颜色、边框和图片尺寸。",
            f"三雷达 Radar-ID：{p('RID_R2D_3CLASS')}。",
            "画图差异去掉以后，仍然几乎一眼就能认出是哪台雷达。",
        ],
        size=20,
    )
    _footer(s, 7)
    _notes(
        s,
        f"我们回到原始 IQ，按各雷达自己的链重算时间–Doppler 图，不再用官方截图。"
        f"去掉画图痕迹之后，三分类雷达识别仍是 {p('RID_R2D_3CLASS')}。"
        "也就是说，雷达身份写在信号里，不只写在 PNG 的边框上。",
    )

    # 8
    s = prs.slides.add_slide(blank)
    _bg(s)
    _title(s, "那把 Doppler 换成真正的速度呢？")
    _bullets(
        s,
        0.55,
        1.3,
        8.4,
        3.8,
        [
            "只比较 24 GHz 和 77 GHz。10 GHz 没有可靠载频，不进这一轴。",
            "同一径向速度，在不同载频下对应的 Doppler 赫兹数不同。",
            f"换到径向速度以后，Radar-ID：{p('RID_R2D_2477')} → {p('RID_R2V_2477')}。",
            "有改善，但远远谈不上对齐。",
        ],
        size=20,
    )
    _text(s, 9.2, 2.4, 3.6, 1.6, "v = c f_D / (2 f_c)", size=16, color=MUTED, align=PP_ALIGN.CENTER)
    _text(s, 9.2, 4.0, 3.6, 0.8, f"24/77 两分类，\n不要和三雷达 {p('RID_R2D_3CLASS')} 混比", size=13, color=MUTED, align=PP_ALIGN.CENTER)
    _footer(s, 8)
    _notes(
        s,
        f"24 和 77 的载频差一截，同样的人体速度会落在不同的赫兹位置。我们把两路都变到米每秒。"
        f"两分类 Radar-ID 从 {p('RID_R2D_2477')} 降到 {p('RID_R2V_2477')}，有一点用，但配对距离几乎没动。"
        f"请不要把这个 {p('RID_R2V_2477')} 和前面的三雷达 {p('RID_R2D_3CLASS')} 画成同一根柱子。10 GHz 进不来，因为有效载频在公开文件里对不上。",
    )

    # 9
    s = prs.slides.add_slide(blank)
    _bg(s)
    _title(s, "如果不是坐标问题，那是不是信号处理流程不同？")
    _text(s, 0.55, 1.2, 12, 0.4, "每排除一层，我们就更接近真正的问题。", size=16, color=MUTED)
    steps = [
        ("PNG", "部分影响"),
        ("数值 Doppler", "身份仍强"),
        ("速度轴", "弱改善"),
        ("统一 STFT", "几乎不变"),
        ("统一 MTI", "更容易认出雷达"),
        ("简单标定", "仍失败"),
    ]
    for i, (a, b) in enumerate(steps):
        x = 0.45 + i * 2.15
        sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(2.3), Inches(2.0), Inches(2.6))
        _set_fill(sh, WHITE)
        sh.line.color.rgb = LINE
        _text(s, x, 2.55, 2.0, 0.9, a, size=15, bold=True, color=NAVY, align=PP_ALIGN.CENTER)
        _text(s, x, 3.6, 2.0, 0.9, b, size=14, color=INK, align=PP_ALIGN.CENTER)
        if i < len(steps) - 1:
            _text(s, x + 1.85, 3.2, 0.4, 0.4, "→", size=18, color=NAVY)
    _footer(s, 9)
    if (fig_dir / "fig9_ladder.png").is_file():
        pass
    _notes(
        s,
        "这一页不要念成六个失败。顺序是：先怀疑图，再算数值 Doppler，再对齐速度，再统一 STFT 和 MTI，再做简单标定。"
        "统一窗长几乎不掉雷达身份；把 24 GHz 的 MTI 接到 77 上，雷达反而更好认。"
        "能排除的简单解释我们都走了一遍。",
    )

    # 10
    s = prs.slides.add_slide(blank)
    _bg(s)
    _title(s, "我们也试过让模型“主动忘掉雷达”")
    rows = [
        ("共享编码器 A0", p("A0_ACC")),
        ("雷达适配器 A1", p("A1_ACC")),
        ("适配器 + 配对对齐 A2", p("A2_ACC")),
        ("最终 Radar-ID", p("RID_A2_EMBED")),
    ]
    for i, (lab, val) in enumerate(rows):
        y = 1.45 + i * 0.85
        _text(s, 0.7, y, 6.5, 0.7, lab, size=20, color=INK)
        _text(s, 7.5, y, 4.5, 0.7, val, size=28, bold=True, color=NAVY if i < 3 else BRICK)
    _text(s, 0.7, 5.0, 12, 1.4, "网络会认动作，但并没有因此忘记雷达。\n强行把不同雷达的特征拉近，动作分类反而变差。", size=18, color=INK)
    _footer(s, 10)
    _notes(
        s,
        f"这是已知三雷达一起训练的结果，不是换一台没见过的雷达。"
        f"共享编码器 {p('A0_ACC')}，加上适配器掉到 {p('A1_ACC')}，再加强配对对齐掉到 {p('A2_ACC')}，嵌入里的雷达识别仍是 {p('RID_A2_EMBED')}。"
        "这是负结果，不要讲成方法创新。",
    )

    # 11
    s = prs.slides.add_slide(blank)
    _bg(s)
    _title(s, "雷达之间的差异不是固定的，它还会随着动作变化")
    left = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.6), Inches(1.7), Inches(5.8), Inches(3.5))
    _set_fill(left, WHITE)
    left.line.color.rgb = LINE
    _text(s, 0.6, 2.0, 5.8, 0.4, "固定校正", size=16, color=MUTED, align=PP_ALIGN.CENTER)
    _text(s, 0.6, 2.6, 5.8, 1.3, f"R² = {d('D2477_B0_R2')}", size=36, bold=True, color=NAVY, align=PP_ALIGN.CENTER)
    _text(s, 0.6, 4.2, 5.8, 0.5, "不管人怎么动，只用一条平均曲线", size=14, color=MUTED, align=PP_ALIGN.CENTER)
    right = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.9), Inches(1.7), Inches(5.8), Inches(3.5))
    _set_fill(right, WHITE)
    right.line.color.rgb = NAVY
    _text(s, 6.9, 2.0, 5.8, 0.4, "加上当前运动状态　24→77", size=16, color=NAVY, align=PP_ALIGN.CENTER)
    _text(s, 6.9, 2.6, 5.8, 1.3, f"R² = {d('D2477_B1_R2')}", size=36, bold=True, color=NAVY, align=PP_ALIGN.CENTER)
    _text(s, 6.9, 4.2, 5.8, 0.5, f"反向 77→24：{d('D7724_B1_R2')}", size=14, color=MUTED, align=PP_ALIGN.CENTER)
    _text(s, 0.6, 5.5, 12, 1.0, "不同雷达之间的差异，和“人在怎么动”是耦合的。这不是一条固定的硬件曲线。", size=18, color=INK)
    _footer(s, 11)
    if (fig_dir / "fig7_dvq.png").is_file():
        pass
    _notes(
        s,
        f"如果 24 和 77 的差别只是固定偏差，一条平均曲线就该够用。平均模型 24 到 77 的 R² 只有 {d('D2477_B0_R2')}。"
        f"把当前运动状态加进去，升到 {d('D2477_B1_R2')}。"
        "注意这个 R² 说的是“两台雷达能量差能不能被预测”，不是动作识别准不准。"
        "发现本身是：差异随动作变，所以不能当成滤镜。",
    )

    # 12
    s = prs.slides.add_slide(blank)
    _bg(s)
    _title(s, "“能预测差异”不等于“能消除差异”")
    _bullets(
        s,
        0.55,
        1.3,
        12.2,
        4.8,
        [
            f"按预测做校正以后，Radar-ID：{p('CAL_RID_BEFORE')} → {p('CAL_RID_AFTER')}。",
            f"{d('K_SELECTED')} 个低秩分量只解释训练差异的 {p('K20_TRAIN_VAR')}。",
            f"测试集上的理想投影：R² = {d('ORACLE_R2_2D')}。",
            "就算系数已经是理想的，还有一大半跨雷达差异解释不了。",
        ],
        size=20,
    )
    _footer(s, 12)
    _notes(
        s,
        f"预测出 D(v) 再往回补，雷达身份从 {p('CAL_RID_BEFORE')} 变成 {p('CAL_RID_AFTER')}，并没有藏住雷达。"
        f"二维差分用 {d('K_SELECTED')} 个分量，训练集只解释 {p('K20_TRAIN_VAR')}，测试集上把真实差分投影到训练基底，R² 也只有 {d('ORACLE_R2_2D')}。"
        "这是上限诊断，不是上线性能。结论是：缺的不只是回归能力。",
    )

    # 13
    s = prs.slides.add_slide(blank)
    _bg(s)
    _title(s, "真正缺的是：人是以什么几何关系被雷达看见的")
    L = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.5), Inches(1.5), Inches(6.0), Inches(4.6))
    _set_fill(L, WHITE)
    L.line.color.rgb = LINE
    _text(s, 0.7, 1.7, 5.6, 0.4, "现在有", size=16, bold=True, color=NAVY)
    _bullets(s, 0.7, 2.2, 5.6, 3.6, ["动作标签", "三雷达时间对齐", "原始信号"], size=18)
    R = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.5), Inches(6.0), Inches(4.6))
    _set_fill(R, RGBColor(0xF8, 0xEB, 0xE3))
    R.line.color.rgb = BRICK
    _text(s, 7.0, 1.7, 5.6, 0.4, "现在没有", size=16, bold=True, color=BRICK)
    _bullets(s, 7.0, 2.2, 5.6, 3.6, ["人–雷达距离", "观察角度 / 朝向", "雷达空间位置", "可恢复的距离轴"], size=18)
    _text(s, 0.5, 6.3, 12.3, 0.6, "这些只是可能重要的隐变量。现有文件里没有它们，所以不能说“几何就是残差的原因”。", size=16, color=INK)
    _footer(s, 13)
    _notes(
        s,
        "公开配对数据里没有距离、方位、朝向，也没有两台雷达对齐的距离轴。"
        "所以几何假设现在无法检验。这一页最容易讲过：缺变量不等于已经证明几何造成了剩下的差。"
        "我们只能说：没有这些记录，运动、雷达、几何、人缠在一起，拆不开。",
    )

    # 14
    s = prs.slides.add_slide(blank)
    _bg(s)
    _title(s, "下一步不是继续调模型，而是重新把实验做对")
    _bullets(
        s,
        0.55,
        1.35,
        12.2,
        4.6,
        [
            "24 GHz 和 77 GHz 同时看。",
            "5–10 人，7 个动作。",
            "距离 2 / 3 / 4 m，角度 0° / 45° / 90°。",
            "记下朝向、原始 IQ、完整雷达配置、视频时间戳。",
            "同一个人、同一个位置、同一个动作，两台雷达一起看。",
        ],
        size=20,
    )
    _footer(s, 14)
    _notes(
        s,
        "采集工具只是把这一页写成可执行的记录规范，里面的演示文件不是论文结果，不要当成已经采到的数据。"
        "设计意图很具体：同一动作、同一几何，换雷达；同一雷达、同一动作，换几何。这样才谈得上拆开传感器、几何和人。",
    )

    # 15
    s = prs.slides.add_slide(blank)
    _bg(s)
    _title(s, "从“做一个分类器”到“知道真正的问题在哪里”")
    stages = [
        ("1  动作识别", "DONE"),
        ("2  跨人员 / 跨雷达对照", "DONE"),
        ("3  跨雷达测量差异诊断", "DONE"),
        ("4  受控多雷达数据", "NEXT"),
    ]
    for i, (lab, stt) in enumerate(stages):
        x = 0.5 + i * 3.2
        sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(1.8), Inches(3.0), Inches(2.5))
        _set_fill(sh, NAVY if stt == "NEXT" else WHITE)
        sh.line.color.rgb = NAVY
        _text(s, x, 2.1, 3.0, 1.1, lab, size=16, bold=True, color=WHITE if stt == "NEXT" else NAVY, align=PP_ALIGN.CENTER)
        _text(s, x, 3.4, 3.0, 0.5, stt, size=18, bold=True, color=RGBColor(0xF3, 0xD3, 0x9E) if stt == "NEXT" else MUTED, align=PP_ALIGN.CENTER)
    _text(s, 0.5, 4.7, 12.3, 1.6, "Learn Motion, Not Sensor\nCurrent status: NOT YET\n现在清楚的是：下一步必须把空间和采集变量记下来。", size=18, color=INK)
    _footer(s, 15)
    _notes(
        s,
        "收束就三句。分类器在同雷达上能做。跨雷达一换就垮。我们把绘图、速度轴、信号处理、适配器都查过了，身份还在。"
        "Learn Motion, Not Sensor 现在是 NOT YET。下一件事是受控采集，不是再堆一层网络。",
    )

    out_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out_path))
    return {"slides": len(prs.slides), "path": str(out_path)}
