"""Write a Desktop Word file with three-line tables and rebuilt figures."""
from __future__ import annotations

from pathlib import Path

from full_manuscript.three_line import add_figure_caption, add_three_line_table


def write_figure_word(manifest: dict, out_path: Path, tables: list[dict] | None = None) -> Path:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Cm, Inches

    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    if tables:
        for spec in tables:
            add_three_line_table(doc, spec)
            doc.add_page_break()

    for item in manifest["figures"]:
        png = Path(item["file"])
        if png.is_file():
            pic = doc.add_paragraph()
            pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
            pic.paragraph_format.first_line_indent = Cm(0)
            run = pic.add_run()
            run.add_picture(str(png), width=Inches(6.3))
        add_figure_caption(doc, item["caption"], size=10)
        doc.add_page_break()

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))
    return out_path
