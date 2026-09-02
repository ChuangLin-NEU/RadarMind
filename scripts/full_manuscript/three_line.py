"""True three-line Word tables. No vertical rules, no full grid, no cell fill."""
from __future__ import annotations

from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


def _font(run, size=9, bold=False, italic=False, name="Times New Roman"):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), name)
    rfonts.set(qn("w:hAnsi"), name)
    rfonts.set(qn("w:eastAsia"), name)
    rfonts.set(qn("w:cs"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor(0x00, 0x00, 0x00)


def _border_el(tag: str, val: str, sz: str) -> OxmlElement:
    el = OxmlElement(f"w:{tag}")
    el.set(qn("w:val"), val)
    el.set(qn("w:sz"), sz)
    el.set(qn("w:space"), "0")
    el.set(qn("w:color"), "000000")
    return el


def _clear_and_set_tc_borders(cell, edges: dict[str, tuple[str, str]]) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    for child in list(tcPr):
        if child.tag == qn("w:tcBorders"):
            tcPr.remove(child)
        if child.tag == qn("w:shd"):
            tcPr.remove(child)
    tcB = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        val, sz = edges.get(edge, ("nil", "0"))
        tcB.append(_border_el(edge, val, sz))
    tcPr.append(tcB)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), "auto")
    tcPr.append(shd)


def apply_three_line(table, n_header: int = 1) -> None:
    tbl = table._tbl
    tblPr = tbl.tblPr if tbl.tblPr is not None else OxmlElement("w:tblPr")
    if tbl.tblPr is None:
        tbl.insert(0, tblPr)
    for child in list(tblPr):
        if child.tag == qn("w:tblBorders"):
            tblPr.remove(child)
        if child.tag == qn("w:tblLook"):
            tblPr.remove(child)
        if child.tag == qn("w:shd"):
            tblPr.remove(child)
    borders = OxmlElement("w:tblBorders")
    # 1.0 pt = 8 eighths; 0.75 pt = 6
    spec = {
        "top": ("single", "8"),
        "left": ("nil", "0"),
        "bottom": ("single", "8"),
        "right": ("nil", "0"),
        "insideH": ("nil", "0"),
        "insideV": ("nil", "0"),
    }
    for edge, (val, sz) in spec.items():
        borders.append(_border_el(edge, val, sz))
    tblPr.append(borders)
    look = OxmlElement("w:tblLook")
    look.set(qn("w:firstRow"), "0")
    look.set(qn("w:lastRow"), "0")
    look.set(qn("w:firstColumn"), "0")
    look.set(qn("w:lastColumn"), "0")
    look.set(qn("w:noHBand"), "1")
    look.set(qn("w:noVBand"), "1")
    tblPr.append(look)

    n_rows = len(table.rows)
    last = n_rows - 1
    for i, row in enumerate(table.rows):
        for cell in row.cells:
            edges = {e: ("nil", "0") for e in ("top", "left", "bottom", "right")}
            if i == 0:
                edges["top"] = ("single", "8")
            if i == n_header - 1:
                edges["bottom"] = ("single", "6")
            if i == last:
                edges["bottom"] = ("single", "8")
            _clear_and_set_tc_borders(cell, edges)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table_caption(doc, title: str, size: int = 9) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    pf = p.paragraph_format
    pf.space_before = Pt(10)
    pf.space_after = Pt(4)
    pf.line_spacing = 1.0
    pf.first_line_indent = Cm(0)
    text = title.strip()
    if ". " in text[:16] and text.split(". ", 1)[0].startswith("Table "):
        prefix, rest = text.split(". ", 1)
        r1 = p.add_run(prefix + ".")
        _font(r1, size=size, bold=True)
        r2 = p.add_run(" " + rest)
        _font(r2, size=size, bold=False)
    else:
        r = p.add_run(text)
        _font(r, size=size, bold=False)


def add_figure_caption(doc, caption: str, size: int = 9) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    pf = p.paragraph_format
    pf.space_before = Pt(4)
    pf.space_after = Pt(10)
    pf.line_spacing = 1.0
    pf.first_line_indent = Cm(0)
    text = caption.strip()
    if ". " in text[:14] and text.split(". ", 1)[0].startswith("Fig. "):
        prefix, rest = text.split(". ", 1)
        r1 = p.add_run(prefix + ".")
        _font(r1, size=size, bold=True)
        r2 = p.add_run(" " + rest)
        _font(r2, size=size, bold=False)
    else:
        r = p.add_run(text)
        _font(r, size=size, bold=False)


def add_three_line_table(doc, spec: dict):
    add_table_caption(doc, spec["title"])
    header = spec["header"]
    rows = spec["rows"]
    numeric_cols = set(spec.get("numeric_cols") or [])
    section_rows = set(spec.get("section_rows") or [])
    group_starts = set(spec.get("group_starts") or [])
    table = doc.add_table(rows=1 + len(rows), cols=len(header))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    try:
        table.style = "Table Normal"
    except (KeyError, ValueError):
        pass

    for j, h in enumerate(header):
        cell = table.rows[0].cells[j]
        cell.text = ""
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT if j not in numeric_cols else WD_ALIGN_PARAGRAPH.RIGHT
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.0
        run = p.add_run(str(h))
        _font(run, size=9, bold=True)

    for i, row in enumerate(rows):
        is_section = i in section_rows
        for j, val in enumerate(row):
            cell = table.rows[i + 1].cells[j]
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(4 if (i in group_starts and i > 0 and j == 0) else 1)
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.line_spacing = 1.0
            if is_section:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            elif j in numeric_cols:
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            run = p.add_run(str(val))
            _font(run, size=9, bold=is_section, italic=False)

        if is_section and len(header) > 1:
            table.rows[i + 1].cells[0].merge(table.rows[i + 1].cells[-1])
            p = table.rows[i + 1].cells[0].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            if p.runs:
                p.runs[0].bold = True

    apply_three_line(table, n_header=1)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(6)
    return table
